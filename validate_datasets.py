#!/usr/bin/env python3
"""Verify all generated CSVs and Excel workbooks, bounds, dimensions and distinctness."""
import csv, hashlib, json
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent

def main():
    index=list(csv.DictReader((ROOT/'dataset_index.csv').open()))
    hashes=set(); checks=0
    for entry in index:
        n=int(entry['hubs']); isn=entry['seed']; folder=ROOT/entry['csv_folder']
        tables={}
        for name,expected,width in [('Eik',n,3),('Njk',n,3),('Bmjk',3*n,4),('Cijk',n*n,4)]:
            with (folder/f'{isn}_Param_{name}.CSV').open(newline='') as f: rows=list(csv.reader(f))
            assert len(rows)==expected+1,(entry,name,len(rows))
            assert all(len(r)==width for r in rows)
            tables[name]=rows
            values=[list(map(int,r)) for r in rows[1:]]
            if name in ('Eik','Njk'):
                assert [r[:2] for r in values]==[[1,j] for j in range(1,n+1)]
                assert all(r[2]==0 or 20<=r[2]<150 for r in values)
            else:
                suppliers=3 if name=='Bmjk' else n
                assert [r[:3] for r in values]==[[1,i,j] for i in range(1,suppliers+1) for j in range(1,n+1)]
                lo,hi=(75,175) if name=='Bmjk' else (50,150)
                assert all(r[3]==0 if r[1]==r[2] else lo<=r[3]<hi for r in values)
        assert all((int(a[2])==0) != (int(b[2])==0) for a,b in zip(tables['Eik'][1:],tables['Njk'][1:]))
        digest=hashlib.sha256(json.dumps(tables,sort_keys=True).encode()).hexdigest()
        assert digest not in hashes,'Duplicate dataset'; hashes.add(digest)
        workbook=load_workbook(ROOT/entry['workbook'],read_only=True)
        assert workbook.sheetnames==['PARA_Eik','PARA_Njk','PARA_Bmjk','PARA_Cijk']
        for name,rows in tables.items():
            sheet=workbook['PARA_'+name]
            count=0
            for expected,actual in zip(rows,sheet.iter_rows(values_only=True),strict=True):
                assert list(actual)==expected
                count+=1
            checks+=count-1
        workbook.close()
    report={'datasets':len(index),'distinct_datasets':len(hashes),'workbooks':len(index),'csv_files':4*len(index),'data_rows_verified_in_both_formats':checks,'status':'PASS'}
    (ROOT/'validation.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report),flush=True)
if __name__=='__main__': main()
